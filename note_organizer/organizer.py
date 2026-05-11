import json
import re
import shutil
import subprocess
from pathlib import Path

from enhanced_tagger import get_tags_for_note

try:
    import docx
except ImportError:
    docx = None

try:
    import spacy
    SPACY_MODEL = "en_core_web_sm"
    try:
        nlp = spacy.load(SPACY_MODEL)
    except Exception:
        nlp = spacy.blank("en")
except ImportError:
    nlp = None

try:
    import openai
except ImportError:
    openai = None

SUPPORTED_EXTENSIONS = [".txt", ".docx", ".doc"]

TAG_PATTERNS = {
    "ethnicity": [
        r"ethnic", r"ethnicity", r"kurd", r"turk", r"alevi", r"y[öo]rük", r"roma", r"circass", r"armen", r"arab",
        r"kürt", r"çerkes", r"göçebe", r"boşnak", r"rum", r"tatar", r"laz",
    ],
    "herding": [
        r"herd", r"sheep", r"cattle", r"goat", r"grazing", r"shepherd", r"flock", r"pastur",
        r"livestock", r"sheepfold", r"fold", r"animal husbandry",
        r"hayvancılık", r"koyun", r"keçi", r"inek", r"sığır", r"sürü",
        r"ağıl", r"mera", r"otlak", r"davar", r"büyükbaş", r"küçükbaş", r"çoban",
    ],
    "agriculture": [
        r"farm", r"farming", r"crop", r"wheat", r"barley", r"soil", r"tractor", r"plow", r"planting", r"harvest",
        r"irrigation", r"field", r"seed", r"agr",
        r"tarım", r"buğday", r"arpa", r"ekim", r"hasat", r"tarla",
        r"sulama", r"çiftçi", r"bahçe", r"pancar", r"traktör",
    ],
    "household": [
        r"cook", r"kitchen", r"dairy", r"cheese", r"butter", r"milk", r"home", r"household", r"family",
        r"\bHH\b",
        r"mutfak", r"peynir", r"tereyağ", r"yoğurt", r"yemek", r"hane",
    ],
    "migration": [
        r"migrat", r"settle", r"move", r"village", r"relocat", r"immigrant", r"seasonal", r"nomad",
        r"göç", r"yerleş", r"taşın", r"iskân",
    ],
    "social": [
        r"married", r"wife", r"husband", r"son", r"daughter", r"visit", r"men", r"women", r"cousin",
        r"relative", r"brother", r"sister", r"family", r"kinship",
        r"aile", r"oğul", r"akraba", r"komşu", r"muhtar", r"kabile", r"torun", r"kadın", r"erkek",
    ],
    "burial": [
        r"burial", r"grave", r"funeral", r"cemeter", r"buried", r"death", r"died", r"tomb",
        r"mezarlık", r"mezar", r"cenaze", r"defin", r"ölüm",
    ],
    "religion": [
        r"mosque", r"church", r"prayer", r"religion", r"islam", r"muslim", r"christian", r"alevi",
        r"holy", r"sacred", r"ritual", r"ceremony",
        r"cami", r"namaz", r"dua", r"müslüman", r"ramazan", r"bayram", r"tekke", r"türbe",
    ],
    "economy": [
        r"tax", r"price", r"market", r"income", r"money", r"cost", r"sell", r"buy", r"trade",
        r"fiyat", r"pazar", r"para", r"gelir", r"vergi", r"ticaret", r"satış",
    ],
    "archaeology": [
        r"excavation", r"excavate", r"trench", r"stratigraphy", r"context", r"layer", r"dig",
        r"survey", r"site", r"artifact", r"pottery", r"lithic", r"ceramic", r"bone", r"architecture",
        r"kazı", r"arkeoloji", r"çanak", r"seramik", r"buluntu", r"tabaka", r"höyük",
    ],
    "artifacts": [
        r"artifact", r"pottery", r"ceramic", r"lithic", r"tool", r"vessel", r"ornament",
        r"bead", r"metal", r"shard", r"bone",
        r"çanak", r"seramik", r"buluntu",
    ],
    "architecture": [
        r"building", r"wall", r"house", r"structure", r"ruin", r"room", r"foundation",
        r"mudbrick", r"stone", r"architecture", r"settlement", r"courtyard",
    ],
    "landscape": [
        r"landscape", r"topography", r"hill", r"valley", r"mound", r"survey", r"plain",
        r"plateau", r"terrace", r"site",
        r"tepe", r"vadi", r"ova", r"dağ", r"nehir", r"höyük",
    ],
    "cultural": [
        r"tradition", r"custom", r"celebration", r"festival", r"ritual", r"ceremony",
        r"dance", r"music", r"song", r"craft",
        r"gelenek", r"görenek", r"düğün", r"türkü",
    ],
}


def is_text_bytes(data: bytes, max_nontext_ratio: float = 0.30) -> bool:
    if not data:
        return False
    if b"\x00" in data:
        return False

    text_chars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x100)))
    nontext = sum(1 for b in data if b not in text_chars)
    return (nontext / len(data)) <= max_nontext_ratio


def read_text_file(path: Path) -> str | None:
    try:
        raw = path.read_bytes()
    except Exception:
        return None

    if not is_text_bytes(raw):
        return None

    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return None


def read_docx(path: Path) -> str:
    if docx is None:
        return ""
    document = docx.Document(path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def read_doc(path: Path) -> str:
    if shutil.which("antiword") is None:
        return ""
    try:
        proc = subprocess.run(["antiword", str(path)], capture_output=True, text=True, check=True)
        return proc.stdout
    except subprocess.CalledProcessError:
        return ""


def infer_year(path: Path) -> str:
    path_parts = [part for part in path.parts if re.fullmatch(r"19\d{2}|20\d{2}", part)]
    if path_parts:
        return path_parts[0]
    match = re.search(r"(19\d{2}|20\d{2})", path.name)
    return match.group(0) if match else "Unknown"


def normalize_text(text: str) -> str:
    return text.replace("\r\n", "\n").strip()


def extract_keyword_tags(text: str) -> set[str]:
    tags = set()
    lowered = text.lower()
    for tag, patterns in TAG_PATTERNS.items():
        for pattern in patterns:
            if re.search(pattern, lowered):
                tags.add(tag)
                break
    return tags


def extract_spacy_tags(text: str, max_tags: int = 8) -> set[str]:
    if nlp is None:
        return set()

    try:
        doc = nlp(text[:10000])
        candidates = {}
        for chunk in doc.noun_chunks:
            token = chunk.text.strip().lower()
            if len(token) < 3 or len(token) > 30:
                continue
            if token.isnumeric():
                continue
            candidates[token] = candidates.get(token, 0) + 1
        sorted_tags = sorted(candidates.items(), key=lambda item: (-item[1], len(item[0])))
        return {tag for tag, _ in sorted_tags[:max_tags]}
    except Exception:
        # spaCy parser/model may not be installed in the deployment environment.
        return set()


def make_excerpt(text: str, max_chars: int = 240) -> str:
    excerpt = text.strip().replace("\n", " ")
    if len(excerpt) <= max_chars:
        return excerpt
    return excerpt[:max_chars].rsplit(" ", 1)[0] + "..."


def extract_openai_tags(text: str, api_key: str, max_tags: int = 8) -> list[str]:
    if openai is None or not api_key:
        return []
    openai.api_key = api_key
    prompt = (
        "Read the following note and return a comma-separated list of 5-10 short topic tags. "
        "Use lowercase words or short phrases. Do not include explanations.\n\n"
        f"Note:\n{text[:3000]}"
    )
    try:
        response = openai.ChatCompletion.create(
            model="gpt-3.5-turbo",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2,
            max_tokens=120,
        )
        raw = response.choices[0].message.content.strip()
        tags = [tag.strip().lower() for tag in re.split(r"[,;\n]+", raw) if tag.strip()]
        return tags[:max_tags]
    except Exception:
        return []


def scan_notes(root: str, index_path: str | Path = "notes_index.json", openai_key: str | None = None, use_openai: bool = False) -> dict:
    root_path = Path(root).expanduser().resolve()
    if not root_path.exists():
        raise FileNotFoundError(f"Root path not found: {root}")
    files = []
    for path in root_path.rglob("*"):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            if path.name.startswith("~$"):
                continue
            files.append(path)

    notes = []
    skipped = 0
    for path in sorted(files):
        if path.suffix.lower() == ".txt":
            text = read_text_file(path)
            if text is None:
                skipped += 1
                continue
        elif path.suffix.lower() == ".docx":
            text = read_docx(path)
        elif path.suffix.lower() == ".doc":
            text = read_doc(path)
        else:
            text = ""

        text = normalize_text(text)
        
        # Use enhanced tagging with filename and path patterns
        tags = get_tags_for_note(path.name, text, path)
        
        # Also try keyword-based tags as fallback
        keyword_tags = extract_keyword_tags(text)
        tags.update(keyword_tags)
        
        # Add spacy tags if available
        tags.update(extract_spacy_tags(text))
        
        if use_openai and openai_key:
            tags.update(extract_openai_tags(text, openai_key))

        note = {
            "title": path.name,
            "path": str(path),
            "year": infer_year(path),
            "tags": sorted(tags),
            "excerpt": make_excerpt(text),
            "source": path.suffix.lower().lstrip('.'),
        }
        notes.append(note)

    index = {
        "root": str(root_path),
        "note_count": len(notes),
        "notes": notes,
    }
    if skipped:
        index["skipped_unreadable_txt"] = skipped
    index_file = Path(index_path)
    index_file.write_text(json.dumps(index, indent=2, ensure_ascii=False))
    return index


def load_index(index_path: str | Path = "notes_index.json") -> dict:
    path = Path(index_path)
    if not path.exists():
        return {"root": "", "note_count": 0, "notes": []}
    return json.loads(path.read_text(encoding="utf-8"))


_MIME_TO_EXT = {
    "text/plain": ".txt",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/msword": ".doc",
    "application/vnd.google-apps.document": ".docx",
}


def scan_notes_from_drive(folder_id: str, drive_client, index_path: str | Path = "notes_index.json") -> dict:
    """List and index all supported files from a Google Drive folder tree."""
    import io as _io

    print(f"Listing files from Google Drive folder {folder_id}...")
    files = drive_client.list_files_recursive(folder_id)
    print(f"Found {len(files)} files in Drive")

    notes = []
    for i, file_info in enumerate(files, start=1):
        name = file_info["name"]
        drive_id = file_info["id"]
        mime_type = file_info["mime_type"]
        drive_path = file_info["drive_path"]

        if name.startswith("~$"):
            continue

        ext = _MIME_TO_EXT.get(mime_type, Path(name).suffix.lower() or ".txt")

        try:
            raw_bytes = drive_client.download_bytes(drive_id, mime_type)

            if ext == ".txt":
                text = None
                for encoding in ("utf-8", "latin-1", "cp1252"):
                    try:
                        text = raw_bytes.decode(encoding)
                        break
                    except UnicodeDecodeError:
                        continue
                text = text or ""
            elif ext == ".docx":
                if docx is not None:
                    try:
                        doc = docx.Document(_io.BytesIO(raw_bytes))
                        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                        text = "\n\n".join(paragraphs)
                    except Exception:
                        text = ""
                else:
                    text = ""
            else:
                text = ""
        except Exception as exc:
            print(f"  Warning: could not read {name}: {exc}")
            text = ""

        text = normalize_text(text)
        path_obj = Path(drive_path)
        tags = get_tags_for_note(name, text, path_obj)
        tags.update(extract_keyword_tags(text))
        tags.update(extract_spacy_tags(text))

        note = {
            "title": name,
            "path": drive_path,
            "drive_id": drive_id,
            "drive_mime_type": mime_type,
            "year": infer_year(path_obj),
            "tags": sorted(tags),
            "excerpt": make_excerpt(text),
            "source": ext.lstrip("."),
        }
        notes.append(note)
        print(f"  [{i}/{len(files)}] {name}")

    index = {
        "root": f"gdrive://{folder_id}",
        "note_count": len(notes),
        "notes": notes,
        "source": "google_drive",
        "folder_id": folder_id,
    }

    try:
        Path(index_path).write_text(json.dumps(index, indent=2, ensure_ascii=False))
    except Exception as exc:
        print(f"Warning: could not cache index to disk: {exc}")

    print(f"Indexed {len(notes)} notes from Drive")
    return index
